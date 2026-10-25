"""Provider-neutral verified XLSX-to-LMSB acquisition used by optional extensions.

The generic engine does not include reference tables or import openpyxl until
an extension requests conversion of its own audited upstream spreadsheet.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import warnings
from collections.abc import Callable, Mapping
from typing import Any, cast

from .compiled import encode_csv
from .exceptions import DataError
from .reference_data import CSV_FIELDS
from .references import (
    fetch_verified_source,
    provider_cache,
    provider_write_lock,
    read_resource,
    save_verified_resource,
)
from .registry import registry
from .tables.table import GrowthTable


def normalize_spreadsheet(source: bytes, table_class: type[GrowthTable]) -> bytes:
    """Normalize a validated source spreadsheet using an extension's column schema."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DataError("XLSX reference conversion requires openpyxl") from exc
    table = cast(Any, table_class).__new__(table_class)
    book = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    try:
        records = book.worksheets[0].iter_rows(values_only=True)
        header = next(records, None)
        if not header or str(header[0]).strip().lower() not in {"month", "day", "height", "length"}:
            raise DataError(f"Unexpected reference spreadsheet header: {header}")
        values: list[list[str]] = []
        for row in records:
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue
            try:
                x = float(row[table.get_col_index_measurement()])
                l = float(row[table.get_col_index_l()])
                m = float(row[table.get_col_index_m()])
                s = float(row[table.get_col_index_s()])
            except (ValueError, TypeError, IndexError) as exc:
                if not isinstance(row[0], (float, int)):
                    continue
                raise DataError(f"Malformed spreadsheet coefficient row: {row!r}") from exc
            factor = table.get_index_factor()
            if factor is not None:
                x *= factor
            values.append(["", repr(x), repr(l), repr(m), repr(s), ""])
        if len(values) < 30:
            raise DataError("Unexpectedly short reference spreadsheet")
        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(CSV_FIELDS)
        writer.writerows(values)
        return output.getvalue().encode("utf-8")
    finally:
        book.close()


def initialize_reference_provider(
    provider_name: str,
    table_classes: Mapping[str, type[GrowthTable]],
    normalize: Callable[[bytes, type[GrowthTable]], bytes],
    filename: str | None = None,
) -> int:
    """Ensure the requested table, or best-effort initialize all missing tables.

    The provider supplies its own normalization logic. Scientific numeric
    coefficients never reside inside the reusable core.
    """
    provider = registry.provider(provider_name)
    manifest = provider.manifest()["tables"]
    directory = provider_cache(provider)
    if filename is not None:
        if not filename.endswith(".lmsb") or filename[:-5] not in manifest:
            raise DataError(f"Unknown {provider_name} reference resource: {filename!r}")
        selected = {filename[:-5]: manifest[filename[:-5]]}
    else:
        selected = manifest
    installed = 0
    with provider_write_lock(directory):
        for name, item in sorted(selected.items()):
            output_name = name + ".lmsb"
            try:
                try:
                    if read_resource(output_name) is not None:
                        continue
                except DataError:
                    if "PEDGROWTH_REFERENCE_DIR" in os.environ or not (directory / output_name).exists():
                        raise
                cls = table_classes.get(name)
                if cls is None:
                    raise DataError(f"Unknown table in audited reference manifest: {name}")
                original = fetch_verified_source(item["source_url"], item["source_sha256"])
                csv_data = normalize(original, cls)
                if hashlib.sha256(csv_data).hexdigest() != item["csv_sha256"]:
                    raise DataError(f"Reference coefficient conversion mismatch: {name}")
                compiled = encode_csv(csv_data)
                if hashlib.sha256(compiled).hexdigest() != item["lmsb_sha256"]:
                    raise DataError(f"Packed reference checksum mismatch: {name}")
                save_verified_resource(directory, output_name, compiled, item["lmsb_sha256"])
                installed += 1
            except (DataError, OSError, ValueError) as exc:
                if filename is not None:
                    raise
                warnings.warn(
                    f"{provider_name} reference {output_name} unavailable: {exc}; "
                    "it will be retried when the table is requested",
                    RuntimeWarning,
                    stacklevel=2,
                )
    if installed:
        from .tables.table import _load_coefficients

        _load_coefficients.cache_clear()
    return installed


def initialize_spreadsheet_provider(
    provider_name: str,
    table_classes: Mapping[str, type[GrowthTable]],
    filename: str | None = None,
) -> int:
    """Use the generic downloader with the WHO-style XLSX normalizer."""
    return initialize_reference_provider(provider_name, table_classes, normalize_spreadsheet, filename)
