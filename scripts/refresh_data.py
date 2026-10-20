"""Verify packaged growth tables against *live official source files*.

Offline integrity:
    python scripts/refresh_data.py --check

Live verification (non-mutating, exit != 0 on any mismatch or network failure):
    python scripts/refresh_data.py --live

Pin an official file's raw SHA-256 ONLY when its normalized content matches:
    python scripts/refresh_data.py --pin-sources

Update coefficient snapshots explicitly after reviewing differences:
    python scripts/refresh_data.py --refresh

The online modes require pip install '.[refresh]'. This script never executes
data as Python code or loads serialized Python objects.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import sys
from collections.abc import Callable, Iterable, Sequence
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any, TypedDict

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pedgrowth.data.cdc.tables import CDC_TABLES
from pedgrowth.reference_data import CDC_BMI_FIELDS, CSV_FIELDS, check_dataset_shape, parse_lms_csv
from pedgrowth.tables.sex import Sex
from pedgrowth.tables.table import GrowthTable

from tests.reference_sets import WHO_TABLES

DATA = Path(__file__).resolve().parents[1] / "validation" / "source_data"
MANIFEST = DATA / "SOURCES.json"
DAY_TO_MONTHS = 1 / 30.4375


def _instantiate(factory: Callable[[], GrowthTable]) -> GrowthTable:
    return factory()


TABLE_CLASSES: dict[str, Callable[[], GrowthTable]] = {
    _instantiate(cls).get_file_name(): cls for cls in [*WHO_TABLES, *CDC_TABLES]
}


class Manifest(TypedDict):
    tables: dict[str, dict[str, Any]]


def digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def check(manifest: Manifest) -> None:
    if set(TABLE_CLASSES) != set(manifest["tables"]):
        raise RuntimeError("Reference manifest and registered table classes do not match")
    for name, info in manifest["tables"].items():
        path = DATA / f"{name}.csv"
        content = path.read_bytes()
        if digest(content) != info["csv_sha256"]:
            raise RuntimeError(f"Bundled reference data SHA-256 mismatch: {path}")
        data = parse_lms_csv(io.StringIO(content.decode("utf-8-sig")), name)
        instance = TABLE_CLASSES[name]()
        check_dataset_shape(
            data,
            sex_specific=instance.get_col_index_sex() is not None,
            extended_bmi=instance.get_col_index_sigma() is not None,
            name=name,
        )
        if sum(len(items) for items in data.values()) != info["rows"]:
            raise RuntimeError(f"Manifest row count mismatch: {name}")
        print(f"OK {name}: {info['rows']} rows")


def fetch(url: str) -> bytes:
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("Online verification requires the 'refresh' extra: requests") from exc
    response = requests.get(url, timeout=20)
    response.raise_for_status()
    content = response.content
    if not content or len(content) > 15_000_000:
        raise RuntimeError(f"Official reference download empty or oversized: {url}")
    return content


def rows_for_table(source: bytes, table_cls: Callable[[], GrowthTable]) -> list[list[str]]:
    """Convert a WHO XLSX or CDC source CSV to the documented normalized schema."""
    table = table_cls()
    name = table.get_file_name()
    if name.startswith("who"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("WHO XLSX conversion requires openpyxl") from exc
        workbook = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
        try:
            xlsx_rows = workbook.worksheets[0].iter_rows(values_only=True)
            header = next(xlsx_rows, None)
            if not header or str(header[0]).strip().lower() not in {"month", "day", "height", "length"}:
                raise RuntimeError(f"Unknown WHO source header in {name}: {header}")
            result = _extract_rows(xlsx_rows, table)
        finally:
            workbook.close()
    else:
        cdc_rows = csv.reader(io.StringIO(source.decode("utf-8-sig")))
        cdc_header = next(cdc_rows, None)
        if not cdc_header or str(cdc_header[0]).strip().lower() != "sex":
            raise RuntimeError(f"Unknown CDC source header in {name}: {cdc_header}")
        result = _extract_rows(cdc_rows, table)
    if len(result) < 30:
        raise RuntimeError(f"Unexpectedly short official reference table: {name}")
    return result


def _extract_rows(rows: Iterable[Sequence[Any]], table: GrowthTable) -> list[list[str]]:
    who = table.get_file_name().startswith("who")
    result = []
    for line, row in enumerate(rows, start=2):
        if not who and row and str(row[0]).strip().lower() == "sex":
            continue  # Some official CDC CSV files repeat the header between sex groups.
        if not row or all(item is None or str(item).strip() == "" for item in row):
            continue
        try:
            index = float(row[table.get_col_index_measurement()])
            l = float(row[table.get_col_index_l()])
            m = float(row[table.get_col_index_m()])
            s = float(row[table.get_col_index_s()])
            sigma_col = table.get_col_index_sigma()
            sigma = float(row[sigma_col]) if sigma_col is not None else None
            p95 = float(row[14]) if table.get_file_name() == "cdc_bmi_for_age" else None
            sex = "" if who else (Sex.Male.value if int(row[0]) == 1 else Sex.Female.value)
        except (TypeError, ValueError, IndexError) as exc:
            # Some WHO worksheets have formatted non-data footers.
            if who and (not isinstance(row[0], (int, float))):
                continue
            raise RuntimeError(f"Malformed official reference row: {table.get_file_name()}:{line}") from exc
        if who and table.get_index_factor() is not None:
            index *= DAY_TO_MONTHS
        normalized = [sex, repr(index), repr(l), repr(m), repr(s), repr(sigma) if sigma is not None else ""]
        if p95 is not None:
            normalized.append(repr(p95))
        result.append(normalized)
    return result


def normalized_source(source: bytes, table_cls: Callable[[], GrowthTable]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(CDC_BMI_FIELDS if table_cls().get_file_name() == "cdc_bmi_for_age" else CSV_FIELDS)
    writer.writerows(rows_for_table(source, table_cls))
    return buffer.getvalue().encode("utf-8")


def live(
    manifest: Manifest,
    *,
    pin: bool = False,
    refresh: bool = False,
    only: set[str] | None = None,
    workers: int = 4,
) -> int:
    """Check every requested file; report network failures rather than ignoring them.

    --pin-sources permits source-hash bookkeeping ONLY after exact normalized
    content equality. --refresh permits updating normalized tables explicitly.
    """
    names = sorted(set(manifest["tables"]) if only is None else only)
    errors = 0
    staged = {}

    def worker(name: str) -> tuple[dict[str, Any], bytes | None]:
        info = manifest["tables"][name]
        try:
            raw = fetch(info["source_url"])
            normalized = normalized_source(raw, TABLE_CLASSES[name])
            expected = (DATA / (name + ".csv")).read_bytes()
            record = {
                "table": name,
                "source_sha256": digest(raw),
                "expected_source_sha256": info.get("source_sha256"),
                "source_hash_match": (digest(raw) == info.get("source_sha256"))
                if info.get("source_sha256")
                else None,
                "normalized_sha256": digest(normalized),
                "expected_normalized_sha256": info["csv_sha256"],
                "content_identical": normalized == expected,
            }
            return record, normalized
        except Exception as exc:
            return {"table": name, "error": f"{type(exc).__name__}: {exc}"}, None

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(worker, name): name for name in names}
        for future in as_completed(futures):
            name = futures[future]
            result, normalized = future.result()
            if "error" in result:
                print(f"FAIL {name} network/parsing: {result['error']}")
                errors += 1
                continue
            hash_ok = result["source_hash_match"] is True
            content_ok = result["content_identical"]
            print(
                f"{'OK' if hash_ok and content_ok else 'CHANGED'} {name}: "
                f"source_sha256={'MATCH' if hash_ok else 'UNPINNED' if result['source_hash_match'] is None else 'DIFF'} "
                f"normalized_content={'IDENTICAL' if content_ok else 'DIFFERENT'} "
                f"source={result['source_sha256']}"
            )
            if not hash_ok or not content_ok:
                errors += 1
            if pin and content_ok:
                manifest["tables"][name]["source_sha256"] = result["source_sha256"]
                errors -= 0 if hash_ok else 1
            if refresh and normalized is not None:
                # A refresh always requires explicit user initiation.
                from pedgrowth.reference_data import parse_lms_csv

                dataset = parse_lms_csv(io.StringIO(normalized.decode()), name)
                cls = TABLE_CLASSES[name]
                instance = cls()
                check_dataset_shape(
                    dataset,
                    sex_specific=instance.get_col_index_sex() is not None,
                    extended_bmi=instance.get_col_index_sigma() is not None,
                    name=name,
                )
                staged[name] = (normalized, result["source_sha256"], sum(map(len, dataset.values())))
    if refresh and len(staged) == len(names):
        for name, (content, sha, rows) in staged.items():
            from pedgrowth.compiled import encode_csv

            manifest["tables"][name]["lmsb_sha256"] = hashlib.sha256(encode_csv(content)).hexdigest()
            (DATA / (name + ".csv")).write_bytes(content)
            manifest["tables"][name].update(csv_sha256=digest(content), source_sha256=sha, rows=rows)
    if (pin or refresh) and (not refresh or len(staged) == len(names)):
        MANIFEST.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")
    if refresh and len(staged) == len(names):
        errors = 0
    print(f"Checked={len(names)}, failures={errors}")
    return 1 if errors else 0


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--check", action="store_true", help="Offline packaged integrity")
    group.add_argument(
        "--live", action="store_true", help="Compare pinned raw SHA and full normalized content"
    )
    group.add_argument("--pin-sources", action="store_true", help="Pin raw SHA only if content is identical")
    group.add_argument("--refresh", action="store_true", help="Explicitly replace all normalized tables")
    parser.add_argument("--table", action="append", help="Limit live operation to named table(s)")
    parser.add_argument("--workers", type=int, default=4)
    args = parser.parse_args()
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    if args.check:
        check(manifest)
        return
    names = set(args.table) if args.table else None
    if names and names - TABLE_CLASSES.keys():
        parser.error(f"Unknown tables: {sorted(names - TABLE_CLASSES.keys())}")
    if not 1 <= args.workers <= 8:
        parser.error("--workers must be between 1 and 8")
    raise SystemExit(
        live(
            manifest,
            pin=args.pin_sources,
            refresh=args.refresh,
            only=names,
            workers=args.workers,
        )
    )


if __name__ == "__main__":
    main()
